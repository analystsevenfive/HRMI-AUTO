"""Build portable HRMI timing workbooks from the authoritative JSON logs."""
import json
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def build(source, output, mode):
    rows = json.loads(Path(source).read_text(encoding="utf-8"))
    reports = mode == "reports"
    book = Workbook()
    sheet = book.active
    sheet.title = "เวลา Export รายงาน" if reports else "เวลาประมวลผล"
    headers = ["รอบการทำงาน", "สาขา"]
    if reports:
        headers.append("ประเภทรายงาน")
    headers += ["ช่วงวันที่", "เวลาเริ่ม", "เวลาสิ้นสุด", "ระยะเวลา", "สถานะ"]
    headers += ["ไฟล์", "รายละเอียด"] if reports else ["รายละเอียด"]
    sheet.append(headers)
    for run_id in dict.fromkeys(row["run_id"] for row in rows):
        group = [r for r in rows if r["run_id"] == run_id]
        successful = {(r["organization"], r.get("report_type")) for r in group if r["status"] == "สำเร็จ"}
        needed = {(org, kind) for org in ("[BB]", "[BK]", "[NTS]", "[PK]", "[SM]")
                  for kind in (("ขาด", "สาย") if reports else (None,))}
        display = list(group)
        if needed <= successful:
            first = min(group, key=lambda r: r["started_at"])
            last = max(group, key=lambda r: r["finished_at"])
            duration = (datetime.fromisoformat(last["finished_at"]) - datetime.fromisoformat(first["started_at"])).total_seconds()
            display.append(dict(first, organization="สรุปรอบ", report_type="ครบ 10 รายงาน" if reports else "",
                                finished_at=last["finished_at"], duration_seconds=duration,
                                status="สำเร็จครบ 5 สาขา", message="", details="", path=""))
        for row in display:
            start = row.get("target_start_date") or row.get("target_date") or ""
            end = row.get("target_end_date") or row.get("target_date") or start
            values = [run_id, row["organization"]]
            if reports:
                values.append(row.get("report_type", ""))
            values += [start if start == end else f"{start} - {end}"]
            values += [datetime.fromisoformat(row[key]).replace(tzinfo=None) for key in ("started_at", "finished_at")]
            values += [float(row.get("duration_seconds", 0)) / 86400, row["status"]]
            values += [row.get("path", ""), row.get("details", "")] if reports else [row.get("message", "")]
            sheet.append(values)
            date_column = 5 if reports else 4
            for column in (date_column, date_column + 1):
                sheet.cell(sheet.max_row, column).number_format = "dd/mm/yyyy hh:mm:ss"
            sheet.cell(sheet.max_row, date_column + 2).number_format = "[h]:mm:ss"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for column in range(1, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 25
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    Path(output).parent.mkdir(parents=True, exist_ok=True)
    book.save(output)


if __name__ == "__main__":
    build(*sys.argv[1:])
